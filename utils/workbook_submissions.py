"""Private, durable workbook staging. No live-data import or approval operations."""
from contextlib import closing, contextmanager
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import uuid
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from utils.backup_lock import file_lock
from utils.importer import clean, column_map, norm, raw_value, utcnow
from utils.workbook_access import submission_countries

MAX_BYTES=16*1024*1024
MAX_UNPACKED=128*1024*1024
MAX_ROWS=100000
MAX_COLUMNS=256
ID=re.compile(r'[0-9a-f]{32}')
SCHEMA="""
CREATE TABLE IF NOT EXISTS workbook_submissions(
    id TEXT PRIMARY KEY,
    uploader_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
    uploader_username TEXT NOT NULL,
    country_id INTEGER REFERENCES countries(id) ON DELETE SET NULL,
    country_name TEXT NOT NULL,
    original_filename TEXT NOT NULL,
    sha256 TEXT,
    size_bytes INTEGER CHECK(size_bytes>=0),
    submission_note TEXT NOT NULL DEFAULT '',
    complete_country INTEGER NOT NULL CHECK(complete_country=1),
    status TEXT NOT NULL CHECK(status IN ('VALIDATING','PENDING_REVIEW','VALIDATION_FAILED','SUBMISSION_FAILED',
        'NEEDS_REVIEW','REJECTED','IMPORTING','IMPORTED','IMPORT_FAILED')),
    validation_summary TEXT NOT NULL DEFAULT '{}',
    feedback TEXT NOT NULL DEFAULT '',
    duplicate_of TEXT REFERENCES workbook_submissions(id) ON DELETE SET NULL,
    review_baseline TEXT,
    reviewer_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
    reviewed_at TEXT,
    rejection_reason TEXT,
    import_id INTEGER REFERENCES imports(id) ON DELETE SET NULL,
    backup_reference TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_workbook_submissions_owner ON workbook_submissions(uploader_id,created_at);
CREATE INDEX IF NOT EXISTS idx_workbook_submissions_queue ON workbook_submissions(status,created_at);
CREATE INDEX IF NOT EXISTS idx_workbook_submissions_duplicate ON workbook_submissions(uploader_id,country_id,sha256);
CREATE TABLE IF NOT EXISTS workbook_submission_events(
    id INTEGER PRIMARY KEY,
    submission_id TEXT NOT NULL REFERENCES workbook_submissions(id),
    actor_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
    event TEXT NOT NULL,
    created_at TEXT NOT NULL
);
"""

class WorkbookValidationError(ValueError):
    """Safe validation feedback containing no stored directory values."""

def _safe_path(path):
    path=Path(path).absolute()
    if any(part.is_symlink() for part in (path,*path.parents)):
        raise ValueError('Submission paths must not contain symlinks.')
    return path

def _connect(project):
    database=_safe_path(Path(project)/'directory.db')
    con=sqlite3.connect(database.as_uri()+'?mode=rw',uri=True)
    con.row_factory=sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    return con

def _eligible(con,user_id,country_id):
    if not any(row['id']==country_id for row in submission_countries(con,user_id)):
        raise PermissionError('Workbook submission is not permitted for this country.')
    return con.execute('SELECT username FROM users WHERE id=?',(user_id,)).fetchone()['username']

def _event(con,submission_id,actor_id,event):
    # A deleted uploader must not prevent preserving the failed submission's history.
    if not con.execute('SELECT 1 FROM users WHERE id=?',(actor_id,)).fetchone(): actor_id=None
    con.execute('INSERT INTO workbook_submission_events(submission_id,actor_id,event,created_at) VALUES(?,?,?,?)',
                (submission_id,actor_id,event,utcnow()))

def _digest(stream):
    value=hashlib.sha256(); size=0
    for chunk in iter(lambda:stream.read(1024*1024),b''):
        value.update(chunk); size+=len(chunk)
    return value.hexdigest(),size

def _validate(path,country_name,country_code):
    try:
        with zipfile.ZipFile(path) as archive:
            entries=archive.infolist()
            if len(entries)>2048 or sum(item.file_size for item in entries)>MAX_UNPACKED:
                raise WorkbookValidationError('Workbook contents exceed the processing limit.')
            if len({item.filename for item in entries})!=len(entries):
                raise WorkbookValidationError('Workbook contains duplicate archive entries.')
            if '[Content_Types].xml' not in archive.namelist() or 'xl/workbook.xml' not in archive.namelist():
                raise WorkbookValidationError('Upload a valid .xlsx workbook.')
        wb=load_workbook(path,read_only=True,data_only=True,keep_links=False)
        try:
            ws=wb.active
            if ws is None: raise WorkbookValidationError('Workbook has no active worksheet.')
            # Do not trust worksheet dimension declarations to allocate unbounded rows/columns.
            ws.reset_dimensions()
            rows=ws.iter_rows(values_only=True)
            headers=next(rows,())
            if not headers or len(headers)>MAX_COLUMNS:
                raise WorkbookValidationError('Workbook needs a header row with at most 256 columns.')
            keys=[norm(value) for value in headers if value is not None and str(value).strip()]
            if len(keys)!=len(set(keys)):
                raise WorkbookValidationError('Workbook has duplicate column headers.')
            mapping=column_map(headers)
            if 'company_name' not in mapping:
                raise WorkbookValidationError('A company-name column is required.')
            accepted={norm(country_name)}|({norm(country_code)} if country_code else set())
            seen=set(); companies=set(); count=0
            ascii_lower=str.maketrans('ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz')
            for row_number,row in enumerate(rows,2):
                if row_number>MAX_ROWS+1: raise WorkbookValidationError('Workbook exceeds 100000 data rows.')
                if len(row)>MAX_COLUMNS: raise WorkbookValidationError('Workbook exceeds 256 columns.')
                if not any(value is not None and str(value).strip() for value in row): continue
                index=mapping['company_name']; company=clean(row[index]) if index<len(row) else None
                if not company: raise WorkbookValidationError(f'Row {row_number}: company name is missing.')
                if 'country' in mapping:
                    index=mapping['country']; value=row[index] if index<len(row) else None
                    if clean(value) and norm(value) not in accepted:
                        raise WorkbookValidationError(f'Row {row_number}: country does not match the selected country.')
                fingerprint=hashlib.sha256(json.dumps([raw_value(value) for value in row],ensure_ascii=False,default=str).encode()).digest()
                if fingerprint in seen:
                    raise WorkbookValidationError(f'Row {row_number}: duplicate workbook row. Review and remove the duplicate.')
                seen.add(fingerprint); companies.add(company.translate(ascii_lower)); count+=1
            if not count: raise WorkbookValidationError('A full-country replacement cannot be empty.')
            return {'worksheet':ws.title,'rows':count,'companies':len(companies),'columns':len(headers),
                    'other_worksheets':len(wb.worksheets)-1,'mode':'replace_country'}
        finally: wb.close()
    except WorkbookValidationError: raise
    except (OSError,ValueError,KeyError,TypeError,zipfile.BadZipFile,EOFError,ParseError):
        raise WorkbookValidationError('Workbook could not be read. Upload a valid, unencrypted .xlsx file.') from None

def create_submission(project,user_id,country_id,filename,stream,note='',complete_country=False):
    """Web callers must authenticate/CSRF-check first. Returns durable status metadata.

    Only submission metadata and private generated files are written. No directory/import tables
    are changed. A staged workbook is immutable through this API and rechecked before use.
    """
    project=_safe_path(project)
    if complete_country is not True: raise ValueError('Confirm this workbook contains the complete country dataset.')
    if (not isinstance(filename,str) or not filename.strip() or len(filename)>255
        or '/' in filename or '\\' in filename or any(ord(ch)<32 for ch in filename)
        or Path(filename).suffix.lower()!='.xlsx'):
        raise ValueError('Choose a plain .xlsx filename without a path.')
    if not isinstance(note,str) or len(note)>2000: raise ValueError('Submission note must be at most 2000 characters.')
    directory=_safe_path(project/'workbook_submissions')
    directory.mkdir(mode=0o700,exist_ok=True)
    submission_id=uuid.uuid4().hex
    partial=directory/('.'+submission_id+'.partial'); destination=directory/(submission_id+'.xlsx')
    lock=project/'.backup-import.lock'
    with file_lock(lock),closing(_connect(project)) as con,con:
        con.execute('BEGIN IMMEDIATE')
        username=_eligible(con,user_id,country_id)
        country=con.execute('SELECT name,code FROM countries WHERE id=?',(country_id,)).fetchone()
        stamp=utcnow()
        con.execute('''INSERT INTO workbook_submissions(id,uploader_id,uploader_username,country_id,country_name,
            original_filename,submission_note,complete_country,status,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,1,'VALIDATING',?,?)''',
            (submission_id,user_id,username,country_id,country['name'],filename,note.strip(),stamp,stamp))
        _event(con,submission_id,user_id,'SUBMITTED')
    owned_partial=False
    try:
        checksum=hashlib.sha256(); size=0
        fd=os.open(partial,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600); owned_partial=True
        with os.fdopen(fd,'wb') as out:
            for chunk in iter(lambda:stream.read(1024*1024),b''):
                size+=len(chunk)
                if size>MAX_BYTES: raise WorkbookValidationError('Workbook exceeds the 16 MB upload limit.')
                checksum.update(chunk); out.write(chunk)
            out.flush(); os.fsync(out.fileno())
        checksum=checksum.hexdigest()
        with file_lock(lock),closing(_connect(project)) as con,con:
            con.execute('BEGIN IMMEDIATE'); _eligible(con,user_id,country_id)
            # Exclusive publication cannot replace a previously submitted workbook.
            os.link(partial,destination)
            duplicate=con.execute('''SELECT id FROM workbook_submissions WHERE uploader_id=? AND country_id=?
                AND sha256=? AND id!=? ORDER BY created_at DESC,id DESC LIMIT 1''',
                (user_id,country_id,checksum,submission_id)).fetchone()
            con.execute('UPDATE workbook_submissions SET sha256=?,size_bytes=?,duplicate_of=?,updated_at=? WHERE id=?',
                        (checksum,size,duplicate['id'] if duplicate else None,utcnow(),submission_id))
        summary=_validate(destination,country['name'],country['code'])
        with file_lock(lock),closing(_connect(project)) as con,con:
            con.execute('BEGIN IMMEDIATE'); _eligible(con,user_id,country_id)
            with destination.open('rb') as incoming:
                if _digest(incoming)!=(checksum,size): raise OSError('Staged workbook changed during validation.')
            con.execute("UPDATE workbook_submissions SET status='PENDING_REVIEW',validation_summary=?,updated_at=? WHERE id=?",
                        (json.dumps(summary),utcnow(),submission_id))
            _event(con,submission_id,user_id,'VALIDATED')
    except Exception as error:
        status='VALIDATION_FAILED' if isinstance(error,WorkbookValidationError) else 'SUBMISSION_FAILED'
        feedback=str(error) if isinstance(error,WorkbookValidationError) else 'Submission did not complete. Please retry or contact an administrator.'
        # Preserve any published blob and its row for diagnosis; never delete earlier submissions.
        with file_lock(lock),closing(_connect(project)) as con,con:
            con.execute('UPDATE workbook_submissions SET status=?,feedback=?,updated_at=? WHERE id=?',
                        (status,feedback,utcnow(),submission_id))
            _event(con,submission_id,user_id,status)
    finally:
        # Only our exclusively-created temporary file is eligible for cleanup.
        if owned_partial: partial.unlink(missing_ok=True)
    with closing(_connect(project)) as con:
        return dict(con.execute('SELECT * FROM workbook_submissions WHERE id=?',(submission_id,)).fetchone())

def _viewer(con,actor_id):
    actor=con.execute('SELECT role,status FROM users WHERE id=?',(actor_id,)).fetchone()
    if not actor or actor['status']!='ACTIVE': raise PermissionError('Submission access denied.')
    if actor['role']=='ADMIN': return True
    if not submission_countries(con,actor_id): raise PermissionError('Submission access denied.')
    return False

def list_submissions(project,actor_id,limit=50,offset=0):
    if type(limit) is not int or not 1<=limit<=100 or type(offset) is not int or offset<0: raise ValueError('Invalid pagination.')
    with closing(_connect(project)) as con:
        admin=_viewer(con,actor_id)
        return [dict(row) for row in con.execute('SELECT * FROM workbook_submissions '+
                ('' if admin else 'WHERE uploader_id=? ')+ 'ORDER BY created_at DESC,id DESC LIMIT ? OFFSET ?',
                (limit,offset) if admin else (actor_id,limit,offset))]

def get_submission(project,actor_id,submission_id):
    """Return authorized metadata without exposing another uploader's history."""
    if not isinstance(submission_id,str) or not ID.fullmatch(submission_id): return None
    with closing(_connect(project)) as con:
        admin=_viewer(con,actor_id)
        row=con.execute('SELECT * FROM workbook_submissions WHERE id=?',(submission_id,)).fetchone()
        return dict(row) if row and (admin or row['uploader_id']==actor_id) else None

@contextmanager
def verified_workbook(project,actor_id,submission_id):
    """Open only an authorized immutable copy; verify this same descriptor before yielding."""
    if not isinstance(submission_id,str) or not ID.fullmatch(submission_id): raise ValueError('Invalid submission ID.')
    with closing(_connect(project)) as con:
        admin=_viewer(con,actor_id)
        row=con.execute('SELECT * FROM workbook_submissions WHERE id=?',(submission_id,)).fetchone()
        if not row or (not admin and row['uploader_id']!=actor_id): raise PermissionError('Submission access denied.')
        if not row['sha256']: raise ValueError('Workbook storage is incomplete.')
    path=_safe_path(Path(project)/'workbook_submissions'/(submission_id+'.xlsx'))
    with path.open('rb') as stream:
        if _digest(stream)!=(row['sha256'],row['size_bytes']): raise ValueError('Workbook checksum verification failed.')
        stream.seek(0); yield stream

def check_snapshot(con,files):
    """Reject a backup that has metadata but lacks its completed, matching private files."""
    if not con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='workbook_submissions'").fetchone(): return
    for submission_id,checksum,size in con.execute('SELECT id,sha256,size_bytes FROM workbook_submissions WHERE sha256 IS NOT NULL'):
        if not ID.fullmatch(submission_id): raise ValueError('Invalid stored submission ID.')
        if files.get('workbook_submissions/'+submission_id+'.xlsx')!={'sha256':checksum,'size':size}:
            raise ValueError('Submission workbook is missing or does not match the database snapshot.')
