"""Mapped Excel import: a workbook whose layout differs from the standard one is read through a JSON
mapping profile (utils/import_profiles/<name>.json), validated, previewed and only then added.

Kept apart from utils/importer.py on purpose. The standard import replaces a country's whole dataset;
this one only adds to it. Existing companies and contacts are left untouched unless the admin chooses
the "update" mode, and even then a blank cell never erases a stored value. Nothing is ever deleted.

Stored rows use the same columns and the same source_data shape as the standard import, so the search,
company page, permissions and overview treat imported records exactly like any other.
"""
import json, re, hashlib, logging, sqlite3
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from utils.importer import ALIASES, norm, clean, raw_value, agent_id_value, agent_id_index, column_map, utcnow
from utils.fields import parse_date, is_empty

PROFILES=Path(__file__).resolve().parent/"import_profiles"
MODES={"add":"Add new companies and new contacts; leave existing records as they are",
       "skip_existing":"Only add companies that are not in the directory yet",
       "update":"Also update existing companies and matching contacts from the file"}
COMPANY_COLUMNS=("agent_id","network","contact_type","address","city","state")
CONTACT_COLUMNS=("name","job_position","email","phone","address","contact_type","landline_no")
FIELDS=set(COMPANY_COLUMNS)|set(CONTACT_COLUMNS)|{"company_name","country"}
PLACEHOLDERS={"-","–","—","n/a","na","tbc","tba"}
EMAIL_RE=re.compile(r"[^@\s,;<>/]+@[^@\s,;<>/]+\.[A-Za-z]{2,}")
MAX_ROWS=20000

class StalePlan(Exception):
    """The file or the directory changed after the preview was shown."""

# ---------- Profiles ----------
def load_profile(name):
    if not re.fullmatch(r"[a-z0-9_]+",name or ""): raise KeyError(name)
    path=PROFILES/f"{name}.json"
    if not path.is_file(): raise KeyError(name)
    profile=json.loads(path.read_text(encoding="utf-8"))
    check_profile(profile)
    return profile
def profiles():
    """Every valid profile; a broken profile file is logged and left out, never breaking the imports page."""
    out=[]
    for path in sorted(PROFILES.glob("*.json")):
        try: out.append(load_profile(path.stem))
        except Exception: logging.exception("Import profile %s is invalid",path.name)
    return out
def profile_for(filename):
    return next((p for p in profiles() if re.match(p["file_pattern"],filename or "",re.I)),None)
def stored_headers(profile):
    return [c["header"] for c in profile["columns"] if c.get("field") or c.get("keep",False)]
def check_profile(p):
    """A profile must map onto the existing system the same way the company page will read it back."""
    cols=p.get("columns") or []
    for c in cols:
        if not c.get("header"): raise ValueError(f"Profile {p.get('name')}: a column has no header")
        if c.get("field") and c["field"] not in FIELDS: raise ValueError(f"Profile {p['name']}: unknown field {c['field']!r}")
    fields=[c["field"] for c in cols if c.get("field")]
    if len(fields)!=len(set(fields)): raise ValueError(f"Profile {p['name']}: a field is mapped twice")
    if not any(c.get("field")=="company_name" and c.get("required") for c in cols):
        raise ValueError(f"Profile {p['name']}: company_name must be mapped and required")
    # The company page and search re-derive each column's meaning from the stored header (importer.ALIASES).
    # A mapped column must come back as its field; a kept-only column must not be mistaken for a database field.
    headers=stored_headers(p); read_back={i:f for f,i in column_map(headers).items()}
    for i,h in enumerate(headers):
        c=next(x for x in cols if x["header"]==h); field=c.get("field")
        if field=="agent_id":
            if agent_id_index(headers)!=i: raise ValueError(f"Profile {p['name']}: {h!r} is not read back as the Agent ID")
        elif field in ALIASES and read_back.get(i)!=field: raise ValueError(f"Profile {p['name']}: {h!r} is not read back as {field}")
        elif not field and i in read_back: raise ValueError(f"Profile {p['name']}: kept column {h!r} would be read back as {read_back[i]}")

# ---------- Helpers ----------
def nocase(s):
    """Company key as the database compares names (trimmed, COLLATE NOCASE: ASCII letters only)."""
    return "".join(ch.lower() if ch.isascii() else ch for ch in str(s).strip())
def loose(s,suffixes):
    """Looser key for spotting the same company spelled differently (case, punctuation, Pvt/Ltd...)."""
    words=re.findall(r"[a-z0-9]+",str(s).lower())
    return " ".join(w for w in words if w not in suffixes)
def emails_in(v): return {e.lower() for e in EMAIL_RE.findall(v or "")}
def digits(v): return re.sub(r"\D","",v or "")
def file_hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def db_signature(con,cid):
    """Changes whenever the country's companies or contacts change."""
    if cid is None: return "none"
    r=con.execute("""SELECT COUNT(*),COALESCE(MAX(co.id),0),COALESCE(MAX(co.updated_at),''),
        (SELECT COUNT(*) FROM contacts ct JOIN companies c2 ON c2.id=ct.company_id WHERE c2.country_id=?),
        (SELECT COALESCE(MAX(ct.id),0) FROM contacts ct JOIN companies c2 ON c2.id=ct.company_id WHERE c2.country_id=?),
        (SELECT COALESCE(MAX(ct.updated_at),'') FROM contacts ct JOIN companies c2 ON c2.id=ct.company_id WHERE c2.country_id=?)
        FROM companies co WHERE co.country_id=?""",(cid,cid,cid,cid)).fetchone()
    return "|".join(map(str,r))
def merge_source(old,new):
    """Stored XLSX row updated with the file's non-empty values; columns the file leaves blank keep theirs."""
    try: rec=json.loads(old) if old else {}
    except ValueError: rec={}
    if not isinstance(rec,dict): rec={}
    for k,v in json.loads(new).items():
        if not is_empty(v) or k not in rec: rec[k]=v
    return json.dumps(rec,ensure_ascii=False)

# ---------- Analysis (read-only) ----------
def analyze(path,db_path,profile,mode="add",fill_down=False,skip_invalid=False):
    """Read, map and validate the workbook and work out what importing would do. Nothing is written."""
    if mode not in MODES: raise ValueError("Unknown import mode")
    issues=[]  # file-level: {"level": "critical"|"warning", "text"}
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb[profile["sheet"]] if profile.get("sheet") in wb.sheetnames else wb.active
        sheet=ws.title; rows=[]
        for r in ws.iter_rows(values_only=True):
            rows.append(r)
            if len(rows)>MAX_ROWS+profile.get("header_row",1): raise ValueError(f"The workbook has more than {MAX_ROWS} rows")
    finally: wb.close()
    hr=profile.get("header_row",1)
    headers=list(rows[hr-1]) if len(rows)>=hr else []
    body=rows[hr:]
    width=max([len(headers)]+[len(r) for r in body])
    headers+= [None]*(width-len(headers))

    # 1. Columns: detect, match to the profile, flag the rest.
    by_norm={}
    for i,h in enumerate(headers):
        if clean(h): by_norm.setdefault(norm(h),[]).append(i)
    for k,idx in by_norm.items():
        if len(idx)>1: issues.append({"level":"critical","text":f"Column “{clean(headers[idx[0]])}” appears {len(idx)} times; the file is ambiguous."})
    index={}; mapping=[]; matched=set()
    for c in profile["columns"]:
        keys=[norm(c["header"])]+[norm(a) for a in c.get("also",[])]
        i=next((by_norm[k][0] for k in keys if k in by_norm),None)
        stored=bool(c.get("field") or c.get("keep",False))
        if i is not None: index[c["header"]]=i; matched.add(i)
        elif c.get("required"): issues.append({"level":"critical","text":f"Required column “{c['header']}” was not found."})
        elif stored: issues.append({"level":"warning","text":f"Column “{c['header']}” was not found; it will be left empty."})
        mapping.append({"excel":clean(headers[i]) if i is not None else None,"letter":get_column_letter(i+1) if i is not None else None,
                        "header":c["header"],"field":c.get("field"),"stored":stored,"required":bool(c.get("required")),
                        "found":i is not None,"review":c.get("review"),"check":c.get("check")})
    for i,h in enumerate(headers):
        if i in matched: continue
        has_data=any(i<len(r) and not is_empty(r[i]) for r in body)
        if not clean(h) and not has_data: continue
        label=clean(h) or f"Column {get_column_letter(i+1)} (no header)"
        mapping.append({"excel":label,"letter":get_column_letter(i+1),"header":None,"field":None,"stored":False,"required":False,"found":True,
                        "review":"Not in the mapping, so it is not imported. Add it to the profile if it should be kept.","check":None,"unknown":True})
        issues.append({"level":"warning","text":f"Column “{label}” is not in the mapping and will not be imported."})
    cols={c["header"]:c for c in profile["columns"]}
    stored=[h for h in stored_headers(profile)]
    stored.sort(key=lambda h: index.get(h,10**6))  # workbook order, as the standard import stores it
    states={norm(s) for s in profile.get("states",[])}
    suffixes=set(profile.get("legal_suffixes",[]))

    # 2. Rows: map and validate.
    out=[]; blank=0; last=None
    for rownum,row in enumerate(body,start=hr+1):
        if not row or all(is_empty(v) for v in row): blank+=1; last=None; continue
        get=lambda h: (row[index[h]] if h in index and index[h]<len(row) else None)
        vals={}
        for h,c in cols.items():
            if c.get("field")=="agent_id": vals["agent_id"]=agent_id_value(get(h))
            elif c.get("field"): vals[c["field"]]=clean(get(h))
        src={h:raw_value(get(h)) for h in stored}
        errs=[]; warns=[]
        if not vals.get("company_name"):
            if fill_down and last:
                vals["company_name"]=last[0]; src[next(h for h,c in cols.items() if c.get("field")=="company_name")]=last[0]
                warns.append(f"No company name; taken from row {last[1]} ({last[0]}).")
            else: errs.append("Company name is missing.")
        for h,c in cols.items():
            v=clean(get(h)); check=c.get("check")
            if not v or not check: continue
            if check=="country" and v.casefold()!=profile["country"].casefold(): errs.append(f"Country is “{v}”, expected {profile['country']}.")
            elif check=="email":
                bad=[p for p in re.split(r"[\s,;/]+",v) if p and not EMAIL_RE.fullmatch(p)]
                if bad: errs.append(f"{h} “{bad[0]}” is not a valid email address.")
            elif check=="phone" and len(digits(v))<6: warns.append(f"{h} “{v}” has fewer than 6 digits.")
            elif check=="date" and v.lower() not in PLACEHOLDERS and parse_date(get(h)) is None:
                warns.append(f"{h} “{v}” is not a readable date; stored as written.")
            elif check=="state" and states and norm(v) not in states: warns.append(f"{h} “{v}” is not a {profile['country']} state or union territory.")
            elif check=="city" and norm(v) in states: warns.append(f"{h} “{v}” is a state name; State and City may be swapped.")
        for h in stored:
            v=src.get(h)
            if isinstance(v,str) and "�" in v: warns.append(f"{h} contains an unreadable character (�): “{v[:60]}”.")
        if vals.get("company_name"): last=(vals["company_name"],rownum)
        out.append({"row":rownum,"vals":vals,"source":json.dumps(src,ensure_ascii=False),"errors":errs,"warnings":warns,
                    "sig":tuple(src.values())})

    # 3. Against the directory: new, existing, duplicate.
    con=sqlite3.connect(f"file:{Path(db_path).as_posix()}?mode=ro",uri=True); con.row_factory=sqlite3.Row
    try:
        crow=con.execute("SELECT id,name FROM countries WHERE name=? COLLATE NOCASE",(profile["country"],)).fetchone()
        cid=crow["id"] if crow else None
        existing={}; loose_existing={}; contacts={}
        if cid is not None:
            for r in con.execute("SELECT id,company_name FROM companies WHERE country_id=?",(cid,)):
                existing[nocase(r["company_name"])]={"id":r["id"],"name":r["company_name"]}
                loose_existing.setdefault(loose(r["company_name"],suffixes),r["company_name"])
            for r in con.execute("SELECT ct.* FROM contacts ct JOIN companies co ON co.id=ct.company_id WHERE co.country_id=? ORDER BY ct.id",(cid,)):
                contacts.setdefault(r["company_id"],[]).append(dict(r))
        signature=db_signature(con,cid)
    finally: con.close()

    def same_person(v,c):
        """Same contact: a shared email, or (without an email) the same name and phone. Addresses are not
        compared: a newer file often corrects the address of the same person."""
        e=emails_in(v.get("email"))
        if e: return bool(e & emails_in(c["email"]))
        return bool(v.get("name")) and nocase(v["name"])==nocase(c["name"] or "") and digits(v.get("phone"))==digits(c["phone"])
    claimed=set()
    def match_contact(v,company_id):
        """The stored contact a row stands for. One person listed at several offices is several rows and
        several stored contacts; each stored contact is matched once, preferring the same address."""
        found=[c for c in contacts.get(company_id,[]) if same_person(v,c)]
        if not found: return None,False
        free=[c for c in found if c["id"] not in claimed]
        best=next((c for c in free if nocase(v.get("address") or "")==nocase(c["address"] or "")),free[0] if free else None)
        if best: claimed.add(best["id"])
        return best,True

    seen_rows={}; seen_email={}; new_companies={}; companies={}; loose_new={}
    for r in out:
        v=r["vals"]
        if r["errors"]: r["action"]="error"; continue
        ck=nocase(v["company_name"]); ex=existing.get(ck)
        key=(ck,r["sig"])
        if key in seen_rows:
            r["action"]="duplicate_file"; r["warnings"].append(f"Identical to row {seen_rows[key]}; skipped."); continue
        seen_rows[key]=r["row"]
        for e in emails_in(v.get("email")):
            if (ck,e) in seen_email: r["warnings"].append(f"Same email as row {seen_email[(ck,e)]} (other details differ); imported as its own contact row."); break
        for e in emails_in(v.get("email")): seen_email.setdefault((ck,e),r["row"])
        if ex:
            r["company_id"]=ex["id"]; r["company"]=ex["name"]
            info=companies.setdefault(ck,{"name":ex["name"],"id":ex["id"],"status":"existing","rows":[],"first":r})
            info["rows"].append(r["row"])
            if mode=="skip_existing": r["action"]="skip_existing"; continue
            match,known=match_contact(v,ex["id"])
            if match: r["contact_id"]=match["id"]; r["action"]="update_contact" if mode=="update" else "duplicate_existing"
            else:
                r["action"]="add_contact"
                # Known person, but every stored contact of theirs already stands for another row: another office.
                if known: r["warnings"].append("Same person as a stored contact, at another office; added as another contact row.")
            if mode=="update" and info["first"] is r: r["update_company"]=True
        else:
            r["company"]=v["company_name"]
            info=companies.setdefault(ck,{"name":v["company_name"],"id":None,"status":"new","rows":[],"first":r})
            info["rows"].append(r["row"])
            if info["first"] is r:
                r["action"]="new_company"; new_companies[ck]=r
                lk=loose(v["company_name"],suffixes)
                if lk in loose_existing: r["warnings"].append(f"Looks like existing company “{loose_existing[lk]}” spelled differently; it will be added as a separate company.")
                elif lk in loose_new: r["warnings"].append(f"Looks like “{loose_new[lk][0]}” (row {loose_new[lk][1]}) spelled differently; both will be added as separate companies.")
                loose_new.setdefault(lk,(v["company_name"],r["row"]))
            else: r["action"]="add_contact"

    # Company-level values that differ between a company's rows: the company record takes its first row,
    # exactly as the standard import does; every contact keeps its own full row.
    for info in companies.values():
        first=info["first"]["vals"]; differ=[]
        for f in ("city","state","network","address"):
            vals={nocase(r["vals"].get(f) or "") for r in out if r.get("company")==info["name"] and r["action"]!="error"}
            if len(vals)>1: differ.append(f.replace("_"," "))
        info["differs"]=differ

    # 4. Summary and whether the import may run.
    count=lambda *a: sum(r["action"] in a for r in out)
    error_rows=count("error")
    critical=[i for i in issues if i["level"]=="critical"]
    blocked=bool(critical) or (error_rows>0 and not skip_invalid)
    summary={"total":len(out),"blank":blank,"valid":len(out)-error_rows,"errors":error_rows,
             "warnings":sum(bool(r["warnings"]) for r in out),
             "duplicates":count("duplicate_file","duplicate_existing"),
             "new_companies":count("new_company"),"new_contacts":count("new_company","add_contact"),
             "update_companies":sum(bool(r.get("update_company")) for r in out),"update_contacts":count("update_contact"),
             "skipped":count("error","duplicate_file","duplicate_existing","skip_existing"),
             "existing_companies":sum(i["status"]=="existing" for i in companies.values())}
    fh=file_hash(path)
    plan=[(r["row"],r["action"],r.get("company_id"),r.get("contact_id"),bool(r.get("update_company"))) for r in out]
    digest=hashlib.sha256(json.dumps([fh,mode,fill_down,skip_invalid,signature,plan,profile["name"]]).encode()).hexdigest()[:20]
    return {"profile":profile["name"],"title":profile["title"],"country":profile["country"],"country_exists":cid is not None,
            "sheet":sheet,"mapping":mapping,"issues":issues,"critical":critical,"rows":out,"companies":companies,
            "summary":summary,"blocked":blocked,"mode":mode,"fill_down":fill_down,"skip_invalid":skip_invalid,
            "hash":digest,"signature":signature}

# ---------- Import (one transaction) ----------
def apply(path,db_path,profile,expected_hash,source_file,mode="add",fill_down=False,skip_invalid=False):
    """Import exactly what the preview showed. Any failure rolls the whole import back."""
    plan=analyze(path,db_path,profile,mode,fill_down,skip_invalid)
    if plan["hash"]!=expected_hash: raise StalePlan("The file or the directory changed after the preview.")
    if plan["blocked"]: raise ValueError("The file has errors that block the import.")
    con=sqlite3.connect(db_path); con.row_factory=sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    now=utcnow(); added=updated_co=updated_ct=new_co=0
    try:
        con.execute("BEGIN IMMEDIATE")
        crow=con.execute("SELECT id FROM countries WHERE name=? COLLATE NOCASE",(profile["country"],)).fetchone()
        # Nothing may have changed between the preview check above and taking the write lock.
        if db_signature(con,crow["id"] if crow else None)!=plan["signature"]: raise StalePlan("The directory changed after the preview.")
        if crow: cid=crow["id"]
        else:
            con.execute("INSERT INTO countries(name,source_file,created_at,updated_at) VALUES(?,?,?,?)",(profile["country"],source_file,now,now))
            cid=con.execute("SELECT last_insert_rowid()").fetchone()[0]
        ids={}
        for r in plan["rows"]:
            a=r["action"]; v=r["vals"]
            if a in ("error","duplicate_file","duplicate_existing","skip_existing"): continue
            ck=nocase(v["company_name"])
            if a=="new_company":
                con.execute("""INSERT INTO companies(country_id,company_name,agent_id,network,contact_type,address,city,state,source_file,source_row,source_data,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(cid,v["company_name"],*(v.get(f) for f in COMPANY_COLUMNS),source_file,r["row"],r["source"],now,now))
                ids[ck]=con.execute("SELECT last_insert_rowid()").fetchone()[0]; new_co+=1
            company_id=r.get("company_id") or ids[ck]
            if r.get("update_company"):
                old=con.execute("SELECT source_data FROM companies WHERE id=?",(company_id,)).fetchone()["source_data"]
                sets=",".join(f"{f}=COALESCE(?,{f})" for f in COMPANY_COLUMNS)
                con.execute(f"UPDATE companies SET {sets},source_data=?,updated_at=? WHERE id=?",
                            (*(v.get(f) for f in COMPANY_COLUMNS),merge_source(old,r["source"]),now,company_id)); updated_co+=1
            if a=="update_contact":
                old=con.execute("SELECT source_data FROM contacts WHERE id=?",(r["contact_id"],)).fetchone()["source_data"]
                sets=",".join(f"{f}=COALESCE(?,{f})" for f in CONTACT_COLUMNS)
                con.execute(f"UPDATE contacts SET {sets},source_data=?,updated_at=? WHERE id=?",
                            (*(v.get(f) for f in CONTACT_COLUMNS),merge_source(old,r["source"]),now,r["contact_id"])); updated_ct+=1
            else:
                con.execute("""INSERT INTO contacts(company_id,name,job_position,email,phone,address,contact_type,landline_no,source_file,source_row,source_data,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(company_id,*(v.get(f) for f in CONTACT_COLUMNS),source_file,r["row"],r["source"],now,now)); added+=1
        s=plan["summary"]
        if (new_co,added,updated_co,updated_ct)!=(s["new_companies"],s["new_contacts"],s["update_companies"],s["update_contacts"]):
            raise RuntimeError("Import counts did not match the preview; nothing was saved.")
        con.commit()
    except Exception:
        con.rollback(); raise
    finally: con.close()
    return {**plan,"result":{"new_companies":new_co,"new_contacts":added,"update_companies":updated_co,"update_contacts":updated_ct}}
