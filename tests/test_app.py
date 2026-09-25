"""End-to-end checks for access control, imports and admin safeguards.

Run from the project folder:  python tests/test_app.py
The project (including directory.db) is copied to a temporary folder first; real data is never touched.
"""
import os, sys, shutil, sqlite3, tempfile, importlib
from pathlib import Path

SRC=Path(__file__).resolve().parent.parent
TMP=Path(tempfile.mkdtemp(prefix="directory-test-"))
for name in ("app.py","utils","templates","static","data","directory.db"):
    p=SRC/name
    if p.is_dir(): shutil.copytree(p,TMP/name,ignore=shutil.ignore_patterns("__pycache__"))
    elif p.exists(): shutil.copy2(p,TMP/name)
sys.path.insert(0,str(TMP)); os.chdir(TMP)
from werkzeug.security import generate_password_hash
from openpyxl import load_workbook
A=importlib.import_module("app")
A.app.config.update(WTF_CSRF_ENABLED=False,TESTING=True)

results=[]
def check(name,cond):
    results.append((name,bool(cond))); print(("PASS " if cond else "FAIL ")+name)
con=sqlite3.connect(A.DB); con.row_factory=sqlite3.Row
q=lambda sql,*a: con.execute(sql,a).fetchone()[0]
PW="Test-Password-2026!"
for u in ("ann","jane","moises","kharla"):
    if not q("SELECT COUNT(*) FROM users WHERE username=?",u):
        con.execute("INSERT INTO users(username,password_hash,role,status) VALUES(?,?,?,?)",(u,"x","USER","ACTIVE"))
con.execute("UPDATE users SET password_hash=? WHERE username<>'kharla'",(generate_password_hash(PW),))
con.execute("UPDATE users SET password_hash=? WHERE username='kharla'",(generate_password_hash("ChangeMe-Kharla-2026!"),))
con.execute("DELETE FROM users WHERE username='fulltest'"); con.commit()
ids={r["username"]:r["id"] for r in con.execute("SELECT id,username FROM users")}
countries=[str(r["id"]) for r in con.execute("SELECT id FROM countries")]
companies=[str(r["id"]) for r in con.execute("SELECT id FROM companies")]
sg=q("SELECT co.id FROM companies co JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore' ORDER BY co.id LIMIT 1")

def client(u,password=PW,ip="127.0.0.1"):
    c=A.app.test_client(); c.environ_base["REMOTE_ADDR"]=ip
    c.post("/login",data={"username":u,"password":password}); return c
adm=client("admin")

# --- Startup cleanup -------------------------------------------------------------------------------
check("startup removed orphaned contacts",q("SELECT COUNT(*) FROM contacts WHERE company_id NOT IN (SELECT id FROM companies)")==0)
check("startup removed orphaned company grants",q("SELECT COUNT(*) FROM user_company_access WHERE company_id NOT IN (SELECT id FROM companies)")==0)
check("secret key persisted for restarts",(TMP/".secret_key").exists() and A.load_secret_key()==A.app.secret_key)

# --- Agent ID search ---------------------------------------------------------------------------------
import re
from utils.importer import agent_id_value
def found(c,term):
    html=c.get("/search",query_string={"q":term}).data.decode()
    return [re.sub("<[^>]+>","",h) for h in re.findall(r"<h3>(.*?)</h3>",html) if "No companies" not in h]
check("startup backfilled agent_id from stored rows",q("SELECT agent_id FROM companies WHERE company_name='Merlion Bridge Logistics Pte. Ltd.'")=="SGP001")
check("agent ID search finds the company",found(adm,"SGP001")==["Merlion Bridge Logistics Pte. Ltd."])
check("agent ID search ignores case and spaces",found(adm,"  sgp001 ")==["Merlion Bridge Logistics Pte. Ltd."])
check("partial agent ID lists every match",len(found(adm,"SGP00"))==q("SELECT COUNT(*) FROM companies WHERE agent_id LIKE 'SGP00%'"))
sg_cid=q("SELECT country_id FROM companies WHERE agent_id='SGP001'")
con.execute("INSERT INTO companies(country_id,company_name,agent_id) VALUES(?,?,?)",(sg_cid,"AAA Agent Prefix Test","SGP0010")); con.commit()
check("exact agent ID ranks before longer IDs",found(adm,"SGP001")[:2]==["Merlion Bridge Logistics Pte. Ltd.","AAA Agent Prefix Test"])
con.execute("DELETE FROM companies WHERE agent_id='SGP0010'"); con.commit()
check("Excel float IDs read as whole numbers",agent_id_value(1001.0)=="1001" and agent_id_value(" SGP001 ")=="SGP001")
check("user without extra field grants can search agent ID",found(client("kharla","ChangeMe-Kharla-2026!"),"SGP001")==["Merlion Bridge Logistics Pte. Ltd."])

# --- Sales column permissions (v4) -------------------------------------------------------------------
adm.post("/admin/users/create",data={"username":"fulltest","password":PW,"role":"FULL_ACCESS"})
ids["fulltest"]=q("SELECT id FROM users WHERE username='fulltest'")
plan={"ann":["x_kgsales"],"jane":["x_pcsales"],"moises":["x_tisales","x_kwsales"]}
for u,s in plan.items():
    adm.post(f"/admin/users/{ids[u]}",data={"role":"USER","status":"ACTIVE","countries":countries,"companies":companies,"sales_fields":s+["x_bogus"]})
check("invalid sales keys rejected",{r[0] for r in con.execute("SELECT field_name FROM user_field_access WHERE user_id=?",(ids["ann"],))}=={"x_kgsales"})
LABELS={"Kargosmart Sales":"Nina","Panda Cargo Sales":"Marco","Tri-Star Logistics Sales":"Ava","KirinWorld Sales":"Leo"}
expect={"ann":{"Kargosmart Sales"},"jane":{"Panda Cargo Sales"},"moises":{"Tri-Star Logistics Sales","KirinWorld Sales"},"fulltest":set(LABELS),"admin":set(LABELS)}
for u,e in expect.items():
    html=client(u).get(f"/company/{sg}").data.decode()
    got={l for l in LABELS if f"<dt>{l}" in html}
    check(f"{u} sees sales columns {sorted(e)}",got==e)
    if u in plan: check(f"{u}: restricted sales values absent from HTML",not [v for l,v in LABELS.items() if l not in e and f">{v}<" in html])
# User company view shows every non-sales column Full Access sees (v7).
import re
def dts(html): return re.findall(r"<dt>([^<]*)",html)
full_dts=dts(client("fulltest").get(f"/company/{sg}").data.decode())
check("user company view: all non-sales columns of Full Access",dts(client("ann").get(f"/company/{sg}").data.decode())==[d for d in full_dts if d.strip() not in set(LABELS)-{"Kargosmart Sales"}])
check("user company view: old sales header kept only as XLSX hint",">KG Sales<" in client("ann").get(f"/company/{sg}").data.decode())

# --- Search only matches visible columns --------------------------------------------------------------
network=q("SELECT network FROM companies WHERE network IS NOT NULL AND network<>'' LIMIT 1")
check("admin can search by network",b"result-card" in adm.get(f"/search?q={network}").data)
check("User without Network field cannot search by it",b"result-card" not in client("ann").get(f"/search?q={network}").data)
check("User can still search visible fields",b"result-card" in client("ann").get("/search?q=Singapore").data)
check("invalid offset does not crash",client("ann").get("/search?q=a&offset=abc").status_code==200)

# --- Search pagination: one card per company, numbered pages ------------------------------------------
import re
india=q("SELECT COUNT(*) FROM companies co JOIN countries cn ON cn.id=co.country_id WHERE cn.name='India'")
seen=[]; page=1
while True:
    h=adm.get(f"/search?q=india&page={page}").data.decode()
    found=re.findall(r'class="result-card" href="/company/(\d+)"',h); seen+=found
    if f'page={page+1}"' not in h or page>20: break
    page+=1
check("search pages cover every matching company exactly once",len(seen)==len(set(seen))==india)
check("summary shows company range",f"of <strong>{india}</strong> companies" in adm.get("/search?q=india").data.decode())
h=adm.get("/search?q=india&page=2").data.decode()
check("page 2 shows 21–40 with current page marked","Showing <strong>21–40</strong>" in h and 'aria-current="page">2<' in h)
check("legacy ?offset=20 opens page 2",'aria-current="page">2<' in adm.get("/search?q=india&offset=20").data.decode())
check("page beyond range clamps to last page",b"result-card" in adm.get("/search?q=india&page=999").data)
h=adm.get("/search?q=%3Cscript%3Ealert(1)%3C/script%3E").data.decode()
check("query is escaped, no script injection","<script>alert(1)" not in h and "No companies match" in h)
check("matched terms highlighted","<mark>" in adm.get("/search?q=india").data.decode())

# --- Admin safeguards ---------------------------------------------------------------------------------
adm.post(f"/admin/users/{ids['admin']}",data={"role":"USER","status":"ACTIVE"})
check("admin cannot demote self",q("SELECT role FROM users WHERE id=?",ids["admin"])=="ADMIN")
adm.post(f"/admin/users/{ids['admin']}",data={"role":"ADMIN","status":"DISABLED"})
check("admin cannot disable self",q("SELECT status FROM users WHERE id=?",ids["admin"])=="ACTIVE")
grants=q("SELECT COUNT(*) FROM user_company_access WHERE user_id=?",ids["ann"])
adm.post(f"/admin/users/{ids['ann']}",data={"role":"FULL_ACCESS","status":"ACTIVE","password":"short"})
check("short password rejects whole save",q("SELECT role FROM users WHERE id=?",ids["ann"])=="USER" and q("SELECT COUNT(*) FROM user_company_access WHERE user_id=?",ids["ann"])==grants)
adm.post(f"/admin/users/{ids['jane']}",data={"role":"USER","status":"DISABLED"})
check("disabled user is locked out",client("jane").get("/search").status_code==302)
check("admin can still demote another admin",adm.post(f"/admin/users/{ids['moises']}",data={"role":"USER","status":"ACTIVE","countries":countries,"companies":companies,"sales_fields":plan["moises"]}).status_code==302)
check("self-delete still blocked",adm.post(f"/admin/users/{ids['admin']}/delete").status_code==302 and q("SELECT COUNT(*) FROM users WHERE id=?",ids["admin"])==1)
log=(TMP/"app.log").read_text(encoding="utf-8",errors="ignore")
check("audit log records who changed what","Admin admin updated user moises" in log and "role ADMIN->USER" in log)

# --- Login throttling ---------------------------------------------------------------------------------
t=A.app.test_client(); t.environ_base["REMOTE_ADDR"]="10.0.0.9"
for _ in range(5): t.post("/login",data={"username":"ann","password":"wrong"})
check("6th attempt throttled even with correct password",t.post("/login",data={"username":"ann","password":PW}).status_code==429)
check("other IPs unaffected",client("ann",ip="10.0.0.10").get("/search").status_code==200)
A.FAILED_LOGINS.clear()

# --- Default password warning -------------------------------------------------------------------------
check("dashboard warns about default passwords",b"Change default passwords" in adm.get("/admin").data)

# --- Re-import keeps company IDs and grants ------------------------------------------------------------
sg_ids={r["id"] for r in con.execute("SELECT co.id FROM companies co JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore'")}
sg_grants=q("SELECT COUNT(*) FROM user_company_access a JOIN companies co ON co.id=a.company_id JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore' AND a.user_id=?",ids["moises"])
sg_contacts=q("SELECT COUNT(*) FROM contacts ct JOIN companies co ON co.id=ct.company_id JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore'")
adm.post("/admin/imports/seed",data={"filename":"singapore.xlsx"})
check("re-import keeps company IDs",{r["id"] for r in con.execute("SELECT co.id FROM companies co JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore'")}==sg_ids)
check("re-import keeps User company grants",q("SELECT COUNT(*) FROM user_company_access a JOIN companies co ON co.id=a.company_id JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore' AND a.user_id=?",ids["moises"])==sg_grants>0)
check("re-import replaces contacts, no duplicates",q("SELECT COUNT(*) FROM contacts ct JOIN companies co ON co.id=ct.company_id JOIN countries cn ON cn.id=co.country_id WHERE cn.name='Singapore'")==sg_contacts)
check("re-import leaves no orphans",q("SELECT COUNT(*) FROM contacts WHERE company_id NOT IN (SELECT id FROM companies)")==0)
check("moises still sees Singapore record after re-import",b"<dt>Tri-Star Logistics Sales" in client("moises").get(f"/company/{sg}").data)
# A workbook without one company: that company, its contacts and grants go; a renamed one is added.
wb=load_workbook(TMP/"data/singapore.xlsx"); ws=wb.active
hdr=[c.value for c in ws[1]]; ci=hdr.index("Company Name Entity")
first=ws.cell(2,ci+1).value
for row in ws.iter_rows(min_row=2):
    if row[ci].value==first: row[ci].value="Brand New Test Co"
wb.save(TMP/"data/singapore.xlsx")
dropped=q("SELECT id FROM companies WHERE company_name=?",first)
adm.post("/admin/imports/seed",data={"filename":"singapore.xlsx"})
check("removed company deleted with its contacts and grants",q("SELECT COUNT(*) FROM companies WHERE id=?",dropped)==0 and q("SELECT COUNT(*) FROM contacts WHERE company_id=?",dropped)==0 and q("SELECT COUNT(*) FROM user_company_access WHERE company_id=?",dropped)==0)
check("new company added",q("SELECT COUNT(*) FROM companies WHERE company_name='Brand New Test Co'")==1)
check("other companies kept their IDs",len(sg_ids-{dropped}-{r["id"] for r in con.execute("SELECT id FROM companies")})==0)

# --- Import failure handling ---------------------------------------------------------------------------
(TMP/"data/broken.xlsx").write_bytes(b"not a workbook")
r=adm.post("/admin/imports/seed",data={"filename":"broken.xlsx"})
check("broken file import fails gracefully and is recorded",r.status_code==302 and q("SELECT status FROM imports ORDER BY id DESC LIMIT 1")=="FAILED")

# --- Admin overview panels -----------------------------------------------------------------------------
import json
from datetime import date, timedelta
def set_source(company_id,**cols):
    row=con.execute("SELECT source_data FROM companies WHERE id=?",(company_id,)).fetchone()
    data=json.loads(row["source_data"] or "{}"); data.update(cols)
    con.execute("UPDATE companies SET source_data=? WHERE id=?",(json.dumps(data),company_id)); con.commit()
fx=[r["id"] for r in con.execute("SELECT id FROM companies ORDER BY id LIMIT 5")]
fx_name={i:q("SELECT company_name FROM companies WHERE id=?",i) for i in fx}
set_source(fx[0],**{"Network Expiry":(date.today()-timedelta(days=10)).strftime("%d-%b-%Y")})
set_source(fx[1],**{"Network Expiry":(date.today()+timedelta(days=30)).strftime("%d-%b-%Y")})
set_source(fx[2],**{"KYC Status":"Pending"})
set_source(fx[3],**{"Agent Status":"Inactive"})
set_source(fx[4],**{"Network Expiry":"see contract"})
att=A.dashboard.attention(con)
listed=lambda key,i: any(x["id"]==i for x in att[key])
check("attention: expired membership listed",listed("expired",fx[0]))
check("attention: membership expiring in 30 days listed",listed("soon",fx[1]) and "30 days" in next(x["detail"] for x in att["soon"] if x["id"]==fx[1]))
check("attention: pending KYC listed",listed("kyc_pending",fx[2]))
check("attention: inactive agent listed",listed("inactive",fx[3]))
check("attention: unreadable date listed, not guessed",listed("unreadable",fx[4]) and not listed("expired",fx[4]))
h=adm.get("/admin").data.decode()
check("attention panel renders with company links","Needs attention" in h and f'href="/company/{fx[0]}"' in h)
multi=[r["company_id"] for r in con.execute("SELECT company_id FROM contacts GROUP BY company_id HAVING COUNT(*)>=2 ORDER BY company_id LIMIT 3")]
con.execute("UPDATE contacts SET email=NULL WHERE company_id=?",(multi[0],))
con.execute("UPDATE contacts SET email=NULL WHERE id=(SELECT MIN(id) FROM contacts WHERE company_id=?)",(multi[1],))
con.execute("UPDATE contacts SET email=' Shared@Example.test ' WHERE id IN ((SELECT MAX(id) FROM contacts WHERE company_id=?),(SELECT MAX(id) FROM contacts WHERE company_id=?))",(multi[1],multi[2])); con.commit()
dq=A.dashboard.data_quality(con)
check("quality: company with no email at all listed",any(x["id"]==multi[0] for x in dq["no_email_companies"]))
check("quality: company with some contacts missing email listed",any(x["id"]==multi[1] for x in dq["missing_email"]))
dup=next((x for x in dq["duplicates"] if x["name"]=="shared@example.test"),None)
check("quality: shared email detected across companies (case/space-insensitive)",dup is not None and "2 companies" in dup["detail"])
check("quality: counts match database",dq["no_email"]==q("SELECT COUNT(*) FROM contacts WHERE email IS NULL OR trim(email)=''") and dq["contacts"]==q("SELECT COUNT(*) FROM contacts"))
check("quality panel renders","Contact data quality" in adm.get("/admin").data.decode())
con.execute("UPDATE companies SET network='TestNetA, TestNetB and TestNetC' WHERE id=?",(fx[0],)); con.execute("UPDATE companies SET network='Andes Net' WHERE id=?",(fx[1],)); con.commit()
bd=A.dashboard.breakdown(con)
check("breakdown: country totals match database",bd["total"]==q("SELECT COUNT(*) FROM companies") and {c["name"]:c["companies"] for c in bd["countries"]}=={r[0]:r[1] for r in con.execute("SELECT cn.name,COUNT(co.id) FROM countries cn LEFT JOIN companies co ON co.country_id=cn.id GROUP BY cn.id")})
nets={n["name"]:n["companies"] for n in bd["networks"]}
check("breakdown: multi-network company counted for each network (comma and \"and\")",nets.get("TestNetA")==nets.get("TestNetB")==nets.get("TestNetC")==1)
check("breakdown: names containing \"and\" are not split",nets.get("Andes Net")==1)
check("breakdown: companies without network counted",bd["no_network"]==q("SELECT COUNT(*) FROM companies WHERE network IS NULL OR trim(network)=''"))
check("breakdown panel links rows to search",'href="/search?q=Singapore"' in adm.get("/admin").data.decode())
set_source(fx[0],**{"TI Sales":None,"KG Sales":"Rep One / Rep Two"})
set_source(fx[1],**{"KG Sales":"rep one"})
sc={s["label"]:s for s in A.dashboard.sales_coverage(con,A.SALES_FIELDS)}
check("sales: one entry per sales column",list(sc)==list(A.SALES_FIELDS.values()))
check("sales: company without TI rep listed as missing",any(x["id"]==fx[0] for x in sc["Tri-Star Logistics Sales"]["missing"]))
kg={r["name"].lower():r["companies"] for r in sc["Kargosmart Sales"]["reps"]}
check("sales: shared cells split and names merged case-insensitively",kg.get("rep one")==2 and kg.get("rep two")==1)
check("sales: assigned + missing = companies",all(s["assigned"]+len(s["missing"])==s["total"]==q("SELECT COUNT(*) FROM companies") for s in sc.values()))
check("sales panel renders","Sales coverage" in adm.get("/admin").data.decode())
con.execute("INSERT INTO users(username,password_hash,role,status) VALUES('noaccess','x','USER','ACTIVE')"); con.commit()
ao={u["username"]:u for u in A.dashboard.access_overview(con,A.SALES_FIELDS)["users"]}
check("access: User visible count requires country AND company grant",ao["moises"]["visible"]==q("SELECT COUNT(*) FROM companies co WHERE EXISTS(SELECT 1 FROM user_company_access a WHERE a.user_id=? AND a.company_id=co.id) AND EXISTS(SELECT 1 FROM user_country_access b WHERE b.user_id=? AND b.country_id=co.country_id)",ids["moises"],ids["moises"]))
check("access: User sales columns shown",ao["moises"]["sales"]==["Tri-Star Logistics Sales","KirinWorld Sales"])
check("access: active User without grants flagged",ao["noaccess"]["no_access"] and not ao["moises"]["no_access"])
h=adm.get("/admin").data.decode()
check("access panel renders and flags no-access users","Access overview" in h and "can't see any records" in h and ">noaccess</a>" in h)
ih=A.dashboard.import_health(con,TMP/"data")
state={f["name"]:f["state"] for f in ih["files"]}
india_uploaded=bool(re.match(r"^[0-9a-f]{12}_",q("SELECT source_file FROM countries WHERE name='India'") or ""))
check("imports: older data/india.xlsx flagged when India was loaded from a newer upload",state.get("india.xlsx")==("stale" if india_uploaded else "ok"))
check("imports: re-imported data/ file not flagged",state.get("singapore.xlsx")=="ok")
check("imports: data/ file for unknown country marked not imported",state.get("broken.xlsx")=="new")
check("imports: failed imports counted",ih["failed"]==q("SELECT COUNT(*) FROM imports WHERE status='FAILED'")>0)
h=adm.get("/admin").data.decode()
check("imports panel renders; FAILED shown with error style","Import health" in h and 'class="status disabled">FAILED' in h and (not india_uploaded or "data/india.xlsx is older" in h))
for _ in range(2): A.app.test_client().post("/login",data={"username":"My-Secret-Pass-Typed-As-Name","password":"x"})
A.app.test_client().post("/login",data={"username":"kharla","password":"wrong"}); A.FAILED_LOGINS.clear()
for h in A.logging.getLogger().handlers: h.flush()
act=A.dashboard.activity(con,TMP/"app.log")
check("activity: admin actions read from app.log, newest first",act["actions"] and act["actions"][0]["text"].startswith("Admin ") and act["actions"][0]["when"]>=act["actions"][-1]["when"])
check("activity: failed sign-ins in last 24h counted",act["failed"]>=3)
tnames=dict(act["targets"])
check("activity: typed non-account usernames never shown",all("secret" not in n.lower() for n in tnames) and tnames.get("unknown username",0)>=2 and tnames.get("kharla",0)>=1)
h=adm.get("/admin").data.decode()
check("activity panel renders without leaking typed text","Recent activity" in h and "My-Secret-Pass" not in h)
check("activity: missing log handled",A.dashboard.activity(con,TMP/"nope.log")["missing"])
import csv, io
con.execute("UPDATE companies SET company_name='=HYPERLINK(\"http://evil.test\",\"x\")' WHERE id=?",(fx[2],))
con.execute("UPDATE contacts SET name='Zoë Ñúñez' WHERE id=(SELECT MIN(id) FROM contacts WHERE company_id=?)",(fx[3],)); con.commit()
r=adm.get("/admin/export/companies.csv")
check("export: companies CSV downloads as attachment",r.status_code==200 and r.mimetype=="text/csv" and "attachment" in r.headers["Content-Disposition"])
rows=list(csv.reader(io.StringIO(r.data.decode("utf-8-sig"))))
check("export: one row per company plus header, XLSX columns included",len(rows)-1==q("SELECT COUNT(*) FROM companies") and "KG Sales" in rows[0] and rows[0][:3]==["Company ID","Country","Company"])
check("export: formula-looking cells neutralized",any(row[2].startswith("'=HYPERLINK") for row in rows[1:]) and not any(row[2].startswith("=") for row in rows[1:]))
r=adm.get("/admin/export/contacts.csv"); crow=list(csv.reader(io.StringIO(r.data.decode("utf-8-sig"))))
check("export: contacts CSV has every contact, UTF-8 names intact",len(crow)-1==q("SELECT COUNT(*) FROM contacts") and any("Zoë Ñúñez" in row for row in crow) and r.data.startswith(b"\xef\xbb\xbf"))
check("export: unknown kind is 404",adm.get("/admin/export/users.csv").status_code==404)
check("export: Users and Full Access cannot export",client("moises").get("/admin/export/contacts.csv").status_code==403 and client("fulltest").get("/admin/export/contacts.csv").status_code==403)
for h in A.logging.getLogger().handlers: h.flush()
check("export: logged with admin name","Admin admin exported contacts" in (TMP/"app.log").read_text(encoding="utf-8",errors="ignore"))
h=adm.get("/admin").data.decode()
qa=re.findall(r'class="quick-action" href="([^"]+)"',h)
check("quick actions: six links, all working for admin",len(qa)==6 and all(adm.get(u).status_code==200 for u in qa))
anchors=re.findall(r'<nav class="section-nav"[^>]*>(.*?)</nav>',h,re.S)[0]
targets=re.findall(r'href="#([^"]+)"',anchors)
check("section nav: every jump link has a matching panel",len(targets)==7 and all(f'id="{t}"' in h for t in targets))
orig=A.dashboard.attention; A.dashboard.attention=lambda c: 1/0
h=adm.get("/admin"); A.dashboard.attention=orig
check("failing panel shows error box, page still loads",h.status_code==200 and b"This panel could not be loaded" in h.data and b"Recent imports" in h.data)

# --- Security review fixes -----------------------------------------------------------------------------
# 1 Log injection: line breaks in typed input stay inside one log line
for brk in ("\n","\r"," ","\x85"):
    A.app.test_client().post("/login",data={"username":f"x{brk}2026-09-24 12:00:00,000 INFO Admin admin deleted user forged","password":"x"})
A.FAILED_LOGINS.clear()
for h in A.logging.getLogger().handlers: h.flush()
check("fix 1: login form cannot forge admin actions in Recent activity",not any("forged" in a["text"] for a in A.dashboard.activity(con,TMP/"app.log")["actions"]))
logtext=(TMP/"app.log").read_text(encoding="utf-8",errors="ignore")
check("fix 1: each attempt stays on one escaped Failed login line",all(f"username=x{e}2026-09-24 12:00:00,000 INFO Admin admin deleted user forged ip=" in logtext for e in ("\\x0a","\\x0d","\\u2028","\\x85")))
# 2/3 Seed accounts: created on first start only; admin edits to kharla are never overridden
kid=q("SELECT id FROM users WHERE username='kharla'")
adm.post(f"/admin/users/{kid}",data={"role":"USER","status":"ACTIVE"})
A.init_db()
check("fix 3: revoked kharla access stays revoked after restart",q("SELECT COUNT(*) FROM user_company_access WHERE user_id=?",kid)==0)
con.execute("UPDATE users SET updated_at=created_at WHERE id=?",(kid,)); con.commit(); A.init_db()
check("fix 3: never-edited kharla still receives all companies",q("SELECT COUNT(*) FROM user_company_access WHERE user_id=?",kid)==q("SELECT COUNT(*) FROM companies"))
adm.post(f"/admin/users/{kid}/delete"); A.init_db()
check("fix 2: deleted kharla is not recreated at restart",q("SELECT COUNT(*) FROM users WHERE username='kharla'")==0)
real_db=A.DB; A.DB=TMP/"fresh.db"; A.init_db()
fresh=sqlite3.connect(A.DB); seeded=sorted(r[0] for r in fresh.execute("SELECT username FROM users")); fresh.close(); A.DB=real_db
check("fix 2: a brand-new database still gets admin and kharla",seeded==["admin","kharla"])
# 4 Password change ends sessions; the admin's own session survives their own change
adm.post("/admin/users/create",data={"username":"sessiontest","password":PW,"role":"USER"})
sid=q("SELECT id FROM users WHERE username='sessiontest'"); victim=client("sessiontest")
check("fix 4: session works before reset",victim.get("/search").status_code==200)
adm.post(f"/admin/users/{sid}",data={"role":"USER","status":"ACTIVE","password":"Another-Pass-2026!"})
check("fix 4: password reset signs the user out everywhere",victim.get("/search").status_code==302)
self_adm=client("admin"); self_adm.post(f"/admin/users/{ids['admin']}",data={"role":"ADMIN","status":"ACTIVE","password":"Admin-New-Pass-2026!"})
check("fix 4: admin changing own password stays signed in",self_adm.get("/admin").status_code==200)
old=client("admin",password="Admin-New-Pass-2026!"); adm=old
# 5 Unknown usernames take as long as real ones
import time as _t
def timed(u):
    c=A.app.test_client(); c.environ_base["REMOTE_ADDR"]=f"10.7.{len(u)}.{sum(map(ord,u))%250}"
    s=_t.perf_counter(); c.post("/login",data={"username":u,"password":"wrong"}); return _t.perf_counter()-s
real=sorted(timed("ann") for _ in range(3))[1]; fake=sorted(timed(f"ghost{i}") for i in range(3))[1]; A.FAILED_LOGINS.clear()
check(f"fix 5: unknown username not measurably faster ({fake*1000:.0f} vs {real*1000:.0f} ms)",fake>real/3)
# 6 One IP guessing many usernames is throttled
spray=A.app.test_client(); spray.environ_base["REMOTE_ADDR"]="10.66.0.1"
codes=[spray.post("/login",data={"username":f"guess{i}","password":"x"}).status_code for i in range(31)]
check("fix 6: 31st attempt from one IP across usernames is throttled",codes[:30].count(429)==0 and codes[30]==429)
check("fix 6: other IPs still sign in",client("ann",ip="10.66.0.2").get("/search").status_code==200)
A.FAILED_LOGINS.clear()
A.CLIENT_IP_HEADER="X-Real-IP"
px=A.app.test_client(); px.environ_base["REMOTE_ADDR"]="10.0.0.254"
for i in range(5): px.post("/login",data={"username":"ann","password":"x"},headers={"X-Real-IP":"203.0.113.9"})
blocked=px.post("/login",data={"username":"ann","password":PW},headers={"X-Real-IP":"203.0.113.9"}).status_code
other=px.post("/login",data={"username":"ann","password":PW},headers={"X-Real-IP":"203.0.113.10"}).status_code
A.CLIENT_IP_HEADER=""; A.FAILED_LOGINS.clear()
check("fix 6: behind a proxy, CLIENT_IP_HEADER throttles per real visitor",blocked==429 and other==302)
# 7 Security headers
r=adm.get("/admin")
check("fix 7: pages refuse to be framed and send safe defaults",r.headers.get("X-Frame-Options")=="DENY" and "frame-ancestors 'none'" in r.headers.get("Content-Security-Policy","") and r.headers.get("X-Content-Type-Options")=="nosniff")
# 8 Long search
r=adm.get("/search?q="+"+".join(f"a{i}" for i in range(3500)))
check("fix 8: very long search returns a page, not an error",r.status_code==200 and b"first 8 words" in r.data)
# 9 Country from re-downloaded file names
from utils.importer import detect_country
from werkzeug.utils import secure_filename
check("fix 9: 'india (1).xlsx' / 'india copy.xlsx' map to India",all(detect_country(secure_filename(n))=="India" for n in ("india (1).xlsx","india copy.xlsx","India (2).xlsx")) and detect_country("sri_lanka.xlsx")=="Sri Lanka")
# 10 Crafted country id
r=adm.post(f"/admin/users/{sid}",data={"role":"USER","status":"ACTIVE","countries":["99999",countries[0]]})
check("fix 10: unknown country id ignored instead of crashing",r.status_code==302 and q("SELECT COUNT(*) FROM user_country_access WHERE user_id=?",sid)==1)
# 11 Idle and absolute session timeouts; pre-upgrade sessions are signed out once
c=client("ann")
with c.session_transaction() as s: s["seen"]=int(_t.time())-A.SESSION_IDLE-5
check("fix 11: idle session expires",c.get("/search").status_code==302)
c=client("ann")
with c.session_transaction() as s: s["started"]=int(_t.time())-A.SESSION_MAX-5
check("fix 11: session expires after the maximum lifetime",c.get("/search").status_code==302)
c=client("ann")
with c.session_transaction() as s: s.pop("pwv"); s.pop("started"); s.pop("seen")
check("fix 11: sessions from before this update are signed out",c.get("/search").status_code==302)
check("fix 11: active session keeps working",client("ann").get("/search").status_code==200)

# --- Existing pages -----------------------------------------------------------------------------------
check("admin pages load",all(adm.get(p).status_code==200 for p in ("/admin","/admin/users","/admin/users/create",f"/admin/users/{ids['ann']}","/admin/imports","/admin/database","/admin/records")))
check("Users blocked from admin",client("ann").get("/admin").status_code==403)

# --- First/last name and search greeting (v7) --------------------------------------------------------
def greeting(u): return re.search(r'class="account-greeting">Welcome, <strong>([^<]*)</strong>',client(u).get("/search").data.decode())
kept={"role":"USER","status":"ACTIVE","countries":countries,"companies":companies,"sales_fields":plan["ann"]}
adm.post(f"/admin/users/{ids['ann']}",data={**kept,"first_name":"  Juan ","last_name":"Dela  Cruz"})
check("names saved trimmed",con.execute("SELECT first_name,last_name FROM users WHERE id=?",(ids["ann"],)).fetchone()[:]==("Juan","Dela Cruz"))
h=adm.get(f"/admin/users/{ids['ann']}").data.decode()
check("edit page shows saved names and display name",'value="Juan"' in h and 'value="Dela Cruz"' in h and "<h1>Juan Dela Cruz</h1>" in h and "@ann" in h)
check("save with names keeps other settings",q("SELECT role FROM users WHERE id=?",ids["ann"])=="USER" and q("SELECT COUNT(*) FROM user_field_access WHERE user_id=?",ids["ann"])==1)
check("header greets by first + last name",greeting("ann") and greeting("ann").group(1)=="Juan Dela Cruz")
adm.post(f"/admin/users/{ids['ann']}",data={**kept,"first_name":"","last_name":"Dela Cruz"})
check("only last name: no stray space",greeting("ann").group(1)=="Dela Cruz")
adm.post(f"/admin/users/{ids['ann']}",data={**kept,"first_name":"x"*61,"last_name":"Y"})
check("too-long name rejects whole save",q("SELECT first_name FROM users WHERE id=?",ids["ann"])=="" and q("SELECT last_name FROM users WHERE id=?",ids["ann"])=="Dela Cruz")
adm.post(f"/admin/users/{ids['ann']}",data={**kept,"first_name":"","last_name":""})
check("no name: greeting falls back to username",greeting("ann").group(1)=="ann")
check("names escaped in greeting",adm.post(f"/admin/users/{ids['ann']}",data={**kept,"first_name":"<b>x</b>"}) and "<b>x</b>" not in client("ann").get("/search").data.decode())
adm.post("/admin/users/create",data={"username":"namedtest","password":PW,"role":"USER","first_name":"Maria","last_name":"Santos"})
check("create user stores names",con.execute("SELECT first_name,last_name FROM users WHERE username='namedtest'").fetchone()[:]==("Maria","Santos"))
check("audit log notes name change","name updated" in (TMP/"app.log").read_text(encoding="utf-8",errors="ignore"))

import hmac, hashlib, json as _json
def hook(body,secret="hook-secret",event="push"):
    raw=_json.dumps(body).encode()
    sig="sha256="+hmac.new(secret.encode(),raw,hashlib.sha256).hexdigest()
    return A.app.test_client().post("/deploy",data=raw,content_type="application/json",headers={"X-Hub-Signature-256":sig,"X-GitHub-Event":event})
check("deploy hook hidden without secret",hook({"ref":"refs/heads/main"}).status_code==404)
A.DEPLOY_SECRET="hook-secret"
check("deploy hook rejects bad signature",hook({"ref":"refs/heads/main"},secret="wrong").status_code==403)
check("deploy hook answers ping",hook({},event="ping").data==b"pong")
check("deploy hook ignores other branches",b"ignored" in hook({"ref":"refs/heads/dev"}).data)
check("deploy hook reports failed pull",hook({"ref":"refs/heads/main"}).status_code==500)  # temp copy is not a git repo
A.DEPLOY_SECRET=""

con.close()
passed=sum(ok for _,ok in results)
print(f"\n{passed}/{len(results)} passed  (temp copy: {TMP})")
sys.exit(0 if passed==len(results) else 1)
