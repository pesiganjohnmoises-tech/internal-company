import re, json, sqlite3, unicodedata
from datetime import datetime, date, timezone
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ALIASES={
"company_name":["company","company name","companyname","company_name","company name entity","agent","organization","organisation"],
"country":["country","nation"],"network":["network","association"],
"contact_type":["contact type","contacttype","contact_type"],
"name":["name","contact name","contactname","person"],
"job_position":["job position","jobposition","position","title","designation"],
"email":["email","e-mail","email address"],"phone":["phone","mobile","telephone","tel"],
"landline_no":["landline no","landline","landline number"],
"city":["city","town"],"state":["state","province","region"],
"address":["address","street","street address","company address","office address","streetaddress"]
}
REQUIRED={"company_name"}
# Agent ID gets its own searchable column but stays out of ALIASES: the detail page shows it from source_data.
AGENT_ID_HEADERS=("agentidfinal","agentid","agentcode")  # norm() of accepted headers
def utcnow(): return datetime.now(timezone.utc).replace(tzinfo=None).isoformat(timespec="seconds")+"Z"
def norm(s):
    s=unicodedata.normalize("NFKD",str(s or "")).encode("ascii","ignore").decode().lower()
    return re.sub(r"[^a-z0-9]","",s)
def detect_country(filename):
    stem=Path(filename).stem.replace("_"," ").replace("-"," ").strip()
    # Browsers save repeat downloads as "india (1).xlsx" / "india copy.xlsx" (secure_filename: "india_1");
    # those suffixes are not part of the country name.
    stem=re.sub(r"(\s*(\(\d+\)|\d+|copy))+$","",stem,flags=re.I).strip() or stem
    return " ".join(w.capitalize() for w in stem.split()) or "Unknown"
def agent_id_value(v):
    """Agent ID as text; Excel stores numeric IDs as floats (1001.0), which must read as "1001"."""
    if isinstance(v,float) and v.is_integer(): v=int(v)
    return clean(v)
def agent_id_index(headers):
    keys=[norm(h) for h in headers]
    return next((keys.index(k) for k in AGENT_ID_HEADERS if k in keys),None)
def agent_id_from_source(raw):
    """Agent ID from a stored source_data row (for rows imported before the agent_id column existed)."""
    try: rec=json.loads(raw) if raw else {}
    except ValueError: return None
    if not isinstance(rec,dict): return None
    by_key={norm(k):v for k,v in rec.items()}
    return next((agent_id_value(by_key[k]) for k in AGENT_ID_HEADERS if k in by_key),None)
def column_map(headers):
    normalized={norm(h):i for i,h in enumerate(headers) if h is not None}
    result={}
    for field,aliases in ALIASES.items():
        for alias in aliases:
            if norm(alias) in normalized:
                result[field]=normalized[norm(alias)]; break
    return result
def clean(v):
    if v is None: return None
    s=str(v).strip()
    return s or None
def header_labels(headers):
    """Display keys for every worksheet column; unnamed headers become 'Column X'."""
    return [clean(h) or f"Column {get_column_letter(i+1)}" for i,h in enumerate(headers)]
def raw_value(v):
    if isinstance(v,(datetime,date)): return v.isoformat()
    if isinstance(v,(int,float,bool)): return v
    return clean(v)
def source_record(labels,row):
    """Complete XLSX row as JSON, in column order, so no column is discarded on import.
    Unnamed columns are kept only when they hold a value."""
    rec={}
    for i,label in enumerate(labels):
        v=raw_value(row[i]) if i<len(row) else None
        if v is None and label.startswith("Column "): continue
        rec[label]=v
    return json.dumps(rec,ensure_ascii=False)
def backfill_source_data(db_path,data_dir):
    """Populate source_data for rows imported before it existed, in place (IDs are kept).
    A row is only filled when the workbook row still carries the same company name."""
    con=sqlite3.connect(db_path); con.row_factory=sqlite3.Row
    try:
        pending=con.execute("SELECT DISTINCT source_file FROM contacts WHERE source_data IS NULL AND source_file IS NOT NULL").fetchall()
        pending+=con.execute("SELECT DISTINCT source_file FROM companies WHERE source_data IS NULL AND source_file IS NOT NULL").fetchall()
        filled=0
        for source_file in {r["source_file"] for r in pending}:
            # Uploaded files are stored as "<12 hex>_<name>.xlsx"; the original lives in data/.
            path=Path(data_dir)/re.sub(r"^[0-9a-f]{12}_","",source_file)
            if not path.is_file(): continue
            wb=load_workbook(path,read_only=True,data_only=True)
            try: rows=list(wb.active.iter_rows(values_only=True))
            finally: wb.close()
            if not rows: continue
            labels=header_labels(rows[0]); name_idx=column_map(rows[0]).get("company_name")
            if name_idx is None: continue
            for table,join in (("contacts","JOIN companies co ON co.id=t.company_id"),("companies","JOIN companies co ON co.id=t.id")):
                for r in con.execute(f"SELECT t.id,t.source_row,co.company_name FROM {table} t {join} WHERE t.source_data IS NULL AND t.source_file=?",(source_file,)).fetchall():
                    if not r["source_row"] or r["source_row"]>len(rows): continue
                    row=rows[r["source_row"]-1]
                    if name_idx<len(row) and (clean(row[name_idx]) or "").lower()==r["company_name"].strip().lower():
                        con.execute(f"UPDATE {table} SET source_data=? WHERE id=?",(source_record(labels,row),r["id"])); filled+=1
        con.commit(); return filled
    finally: con.close()
def import_workbook(path,db_path,country,replace_country=False):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb.active; rows=ws.iter_rows(values_only=True)
        try: headers=next(rows)
        except StopIteration: raise ValueError("Workbook is empty")
        mapping=column_map(headers); labels=header_labels(headers); agent_col=agent_id_index(headers)
        missing=REQUIRED-set(mapping)
        if missing: raise ValueError("Missing required column(s): "+", ".join(sorted(missing)))
        data=[]; errors=[]
        for rownum,row in enumerate(rows,start=2):
            if not row or not any(v is not None and str(v).strip() for v in row): continue
            vals={f:(clean(row[i]) if i is not None and i<len(row) else None) for f,i in mapping.items()}
            if not vals.get("company_name"):
                errors.append(f"Row {rownum}: missing company name"); continue
            vals["agent_id"]=agent_id_value(row[agent_col]) if agent_col is not None and agent_col<len(row) else None
            vals["source_data"]=source_record(labels,row)
            data.append((rownum,vals))
    finally: wb.close()
    effective_country=country or detect_country(Path(path).name)
    con=sqlite3.connect(db_path); con.row_factory=sqlite3.Row
    # Must be set outside the transaction; without it deletes leave orphaned contacts and access rows.
    con.execute("PRAGMA foreign_keys=ON")
    try:
        now=utcnow()
        con.execute("BEGIN IMMEDIATE")
        con.execute("INSERT OR IGNORE INTO countries(name,source_file,created_at,updated_at) VALUES(?,?,?,?)",(effective_country,Path(path).name,now,now))
        crow=con.execute("SELECT id FROM countries WHERE name=? COLLATE NOCASE",(effective_country,)).fetchone()
        cid=crow["id"]
        key=lambda name: con.execute("SELECT lower(trim(?))",(name,)).fetchone()[0]
        kept={}
        if replace_country:
            # Companies still in the workbook are refreshed in place (matched by name) so their IDs and
            # user access grants survive; the rest are deleted, cascading their contacts and access rows.
            new_keys={key(v["company_name"]) for _,v in data}
            for r in con.execute("SELECT id,lower(trim(company_name)) k FROM companies WHERE country_id=? ORDER BY id",(cid,)).fetchall():
                if r["k"] in new_keys and r["k"] not in kept: kept[r["k"]]=r["id"]
            con.execute("DELETE FROM contacts WHERE company_id IN (SELECT id FROM companies WHERE country_id=?)",(cid,))
            stale=[(r["id"],) for r in con.execute("SELECT id FROM companies WHERE country_id=?",(cid,)).fetchall() if r["id"] not in kept.values()]
            con.executemany("DELETE FROM companies WHERE id=?",stale)
        inserted=0
        for rownum,v in data:
            company=v["company_name"]
            if kept and key(company) in kept:
                # First row of a company that already existed: same values a fresh insert would get.
                company_id=kept.pop(key(company))
                con.execute("""UPDATE companies SET company_name=?,agent_id=?,network=?,contact_type=?,address=?,city=?,state=?,source_file=?,source_row=?,source_data=?,updated_at=? WHERE id=?""",
                    (company,v.get("agent_id"),v.get("network"),v.get("contact_type"),v.get("address"),v.get("city"),v.get("state"),Path(path).name,rownum,v["source_data"],now,company_id))
                con.execute("""INSERT INTO contacts(company_id,name,job_position,email,phone,address,contact_type,landline_no,source_file,source_row,source_data,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(company_id,v.get("name"),v.get("job_position"),v.get("email"),v.get("phone"),v.get("address"),v.get("contact_type"),v.get("landline_no"),Path(path).name,rownum,v["source_data"],now,now))
                inserted+=1; continue
            existing=con.execute("SELECT id FROM companies WHERE country_id=? AND lower(trim(company_name))=lower(trim(?))",(cid,company)).fetchone()
            if existing: company_id=existing["id"]
            else:
                con.execute("""INSERT INTO companies(country_id,company_name,agent_id,network,contact_type,address,city,state,source_file,source_row,source_data,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(cid,company,v.get("agent_id"),v.get("network"),v.get("contact_type"),v.get("address"),v.get("city"),v.get("state"),Path(path).name,rownum,v["source_data"],now,now))
                company_id=con.execute("SELECT last_insert_rowid()").fetchone()[0]
            # One contact row per source row; no deduplication.
            con.execute("""INSERT INTO contacts(company_id,name,job_position,email,phone,address,contact_type,landline_no,source_file,source_row,source_data,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",(company_id,v.get("name"),v.get("job_position"),v.get("email"),v.get("phone"),v.get("address"),v.get("contact_type"),v.get("landline_no"),Path(path).name,rownum,v["source_data"],now,now))
            inserted+=1
        con.execute("UPDATE countries SET source_file=?,updated_at=? WHERE id=?",(Path(path).name,now,cid))
        con.commit()
    except Exception:
        con.rollback(); raise
    finally: con.close()
    return {"processed":len(data)+len(errors),"imported":inserted,"duplicates":0,"errors":errors}
